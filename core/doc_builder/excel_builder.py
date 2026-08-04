"""Constructor de tablas de Excel con diseño (rápido, con openpyxl).

Toma un TablePlan y arma un .xlsx bonito de una sola pasada: título, encabezados
con color, filas alternas suaves, bordes finos, anchos automáticos, panel fijo,
autofiltro y, si se pide, una fila de TOTAL con fórmula SUMA.

No usa COM ni abre Excel: escribe el archivo directo. Por eso es rápido incluso
con cientos de filas (lo lento de antes era ir celda por celda por COM).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .plan import TablePlan, get_palette

_FONT = "Calibri"  # fuente profesional y legible


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def build_table(plan: TablePlan, out_path: str | Path) -> str:
    """Construye el .xlsx a partir del plan y lo guarda. Devuelve la ruta."""
    plan = plan.validate()
    pal = get_palette(plan.palette)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.sheet_name = plan.sheet_name
    ws.title = plan.sheet_name

    ncols = len(plan.headers)
    borde = Border(*(Side(style="thin", color="C8C8C8"),) * 4)

    # ---- Título (fila 1, fusionada) --------------------------------------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=plan.title)
    c.font = Font(name=_FONT, size=16, bold=True, color=pal.dark)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    fila = 2
    # ---- Subtítulo (opcional) --------------------------------------------
    if plan.subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        c = ws.cell(row=2, column=1, value=plan.subtitle)
        c.font = Font(name=_FONT, size=11, italic=True, color="666666")
        c.alignment = Alignment(horizontal="left", vertical="center")
        fila = 3

    header_row = fila
    # ---- Encabezados ------------------------------------------------------
    for j, titulo in enumerate(plan.headers, start=1):
        c = ws.cell(row=header_row, column=j, value=str(titulo))
        c.font = Font(name=_FONT, size=11, bold=True, color="FFFFFF")
        c.fill = _fill(pal.primary)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = borde
    ws.row_dimensions[header_row].height = 22

    # ---- Datos (con banda alterna) ---------------------------------------
    primera_datos = header_row + 1
    for i, r in enumerate(plan.rows):
        fila_actual = primera_datos + i
        banda = (i % 2 == 1)
        for j, val in enumerate(r, start=1):
            c = ws.cell(row=fila_actual, column=j, value=val)
            c.font = Font(name=_FONT, size=11, color=pal.text)
            c.alignment = Alignment(
                horizontal="center" if _es_numero(val) else "left",
                vertical="center")
            c.border = borde
            if banda:
                c.fill = _fill(pal.band)

    ultima_datos = primera_datos + len(plan.rows) - 1 if plan.rows else header_row

    # ---- Fila de totales (opcional, con fórmula SUMA) --------------------
    if plan.total_columns and plan.rows:
        fila_total = ultima_datos + 1
        etiqueta = ws.cell(row=fila_total, column=1, value="TOTAL")
        etiqueta.font = Font(name=_FONT, size=11, bold=True, color=pal.dark)
        etiqueta.fill = _fill(pal.band)
        etiqueta.border = borde
        for j in range(1, ncols + 1):
            col_idx0 = j - 1
            celda = ws.cell(row=fila_total, column=j)
            celda.border = borde
            celda.fill = _fill(pal.band)
            if col_idx0 in plan.total_columns and j != 1:
                letra = get_column_letter(j)
                celda.value = f"=SUM({letra}{primera_datos}:{letra}{ultima_datos})"
                celda.font = Font(name=_FONT, size=11, bold=True, color=pal.dark)
                celda.alignment = Alignment(horizontal="center", vertical="center")

    # ---- Nota al pie (opcional) ------------------------------------------
    if plan.note:
        fila_nota = (ultima_datos + 3) if plan.total_columns else (ultima_datos + 2)
        ws.merge_cells(start_row=fila_nota, start_column=1,
                       end_row=fila_nota, end_column=ncols)
        c = ws.cell(row=fila_nota, column=1, value=plan.note)
        c.font = Font(name=_FONT, size=9, italic=True, color="888888")

    # ---- Anchos automáticos por contenido --------------------------------
    for j in range(1, ncols + 1):
        largo = len(str(plan.headers[j - 1]))
        for r in plan.rows:
            largo = max(largo, len(str(r[j - 1])))
        ws.column_dimensions[get_column_letter(j)].width = min(48, max(10, largo + 3))

    # ---- Panel fijo (encabezado + título visibles al desplazar) ---------
    ws.freeze_panes = ws.cell(row=primera_datos, column=1)

    # ---- Autofiltro sobre los encabezados --------------------------------
    if plan.rows:
        ws.auto_filter.ref = (
            f"{get_column_letter(1)}{header_row}:"
            f"{get_column_letter(ncols)}{ultima_datos}"
        )

    wb.save(out)
    return str(out)


def _es_numero(v) -> bool:
    if isinstance(v, (int, float)):
        return True
    try:
        float(str(v).replace(",", "").replace("%", "").strip())
        return True
    except (TypeError, ValueError):
        return False
