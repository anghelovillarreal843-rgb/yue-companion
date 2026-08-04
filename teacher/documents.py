"""Lectura de fuentes para el Modo Profesora: documentos, web y OCR.

Cada lector es OPCIONAL y degradable: si falta la librería, no rompe nada;
devuelve un aviso claro de qué instalar. Formatos contemplados:

    .txt .md        texto plano (siempre)
    .pdf            pdfplumber -> pypdf/PyPDF2 (si hay alguno)
    .docx           python-docx
    .pptx           python-pptx
    .epub           ebooklib (+ limpieza de HTML) o respaldo por ZIP
    imágenes        OCR: core.screen_text -> pytesseract (si hay alguno)
    http(s)://…     descarga y limpieza básica de HTML

Así, al activar el Teacher Mode, YUE puede "leer" lo que el usuario le pase sin
depender de un stack fijo. Todo esto es material de apoyo para enseñar; no altera
la memoria ni la personalidad de YUE.
"""
from __future__ import annotations

import os
import re
import zipfile
from dataclasses import dataclass

# ADITIVO: saneador de rutas de Windows dañadas (falta de backslash tras la
# unidad y texto pegado tras la extensión). Ver teacher/path_repair.py.
try:
    from . import path_repair
except Exception:  # pragma: no cover - respaldo si se importa suelto
    import path_repair  # type: ignore


URL_RE = re.compile(r"(https?://[^\s]+)", re.IGNORECASE)
# NUEVO: enlaces locales tipo file:///C:/Users/.../archivo.pdf (con %20, etc.)
FILE_URL_RE = re.compile(r"(file://[^\s\"']+)", re.IGNORECASE)
PATH_HINT_RE = re.compile(r"[\"']?([A-Za-z]:\\[^\"']+|/[^\s\"']+|[^\s\"']+\.(?:pdf|docx|pptx|epub|txt|md|png|jpg|jpeg))[\"']?", re.IGNORECASE)


@dataclass
class SourceText:
    ok: bool
    text: str
    kind: str          # "txt" | "pdf" | "docx" | "pptx" | "epub" | "web" | "ocr" | "error"
    origin: str        # ruta o URL
    note: str = ""     # aviso (p. ej. "instala pdfplumber")


def find_source(text: str) -> str | None:
    """Detecta si el mensaje trae una URL (http o file://) o una ruta de archivo."""
    m = URL_RE.search(text or "")
    if m:
        return m.group(1).rstrip(").,")
    # NUEVO: enlaces locales file:///... (los usa el navegador al abrir un PDF).
    m = FILE_URL_RE.search(text or "")
    if m:
        return m.group(1).rstrip(").,")
    m = PATH_HINT_RE.search(text or "")
    if m:
        cand = m.group(1)
        # ADITIVO: la ruta puede llegar dañada (falta el backslash tras la unidad,
        # p. ej. 'C:Users\\...') o con texto pegado tras la extensión
        # ('...x.pdf por favor'). Probamos variantes saneadas y nos quedamos con
        # la primera que EXISTA en disco.
        for c in path_repair.candidates(cand):
            if os.path.exists(c):
                return c
        # Aunque ninguna exista aún, si hay una con extensión de documento la
        # devolvemos ya reparada (mejor una ruta absoluta correcta que la rota).
        reparada = path_repair.repair_windows_path(path_repair.trim_to_document(cand))
        if os.path.splitext(reparada)[1].lower() in (".pdf", ".docx", ".pptx", ".epub", ".txt", ".md"):
            return reparada
        # Respaldo al comportamiento original por si algo no encajó arriba.
        if os.path.splitext(cand)[1].lower() in (".pdf", ".docx", ".pptx", ".epub", ".txt", ".md"):
            return cand
    return None


def to_local_path(path_or_url: str) -> str:
    """Convierte un enlace file:///C:/...%20... en una ruta real del sistema.

    - Decodifica %20 y demás.
    - En Windows arregla '/C:/Users/...' -> 'C:/Users/...'.
    - Si no es un file://, devuelve la cadena tal cual (quitando comillas).
    """
    s = (path_or_url or "").strip().strip('"\'')
    if s.lower().startswith("file:"):
        from urllib.parse import urlparse, unquote
        p = urlparse(s)
        ruta = unquote(p.path or "")
        # host en UNC (\\servidor\...) si viniera en netloc
        if p.netloc and p.netloc.lower() not in ("", "localhost"):
            ruta = "//" + p.netloc + ruta
        # Windows: '/C:/Users/...' -> 'C:/Users/...'
        if re.match(r"^/[A-Za-z]:", ruta):
            ruta = ruta[1:]
        return ruta
    # ADITIVO: no es file://; saneamos rutas de Windows a las que les falta el
    # backslash tras la unidad ('C:Users\\...' -> 'C:\\Users\\...') para que
    # os.path.exists y el visor las encuentren.
    return path_repair.repair_windows_path(s)


def _ensure_local_file(path_or_url: str) -> str | None:
    """Devuelve una ruta local usable: resuelve file:// y descarga http(s) a un
    temporal. Devuelve None si no se puede obtener."""
    s = to_local_path(path_or_url)
    low = s.lower()
    if low.startswith(("http://", "https://")):
        try:
            import tempfile
            import urllib.request
            ext = os.path.splitext(low.split("?")[0])[1] or ".bin"
            fd, tmp = tempfile.mkstemp(suffix=ext)
            os.close(fd)
            req = urllib.request.Request(s, headers={"User-Agent": "YUE-Teacher/1.0"})
            with urllib.request.urlopen(req, timeout=20) as r, open(tmp, "wb") as f:
                f.write(r.read())
            return tmp
        except Exception as exc:
            print("[profesora] no pude descargar el enlace:", exc)
            return None
    return s if os.path.exists(s) else None


def read_source(path_or_url: str, max_chars: int = 12000) -> SourceText:
    """Lee una fuente y devuelve su texto (recortado a `max_chars`)."""
    origin = (path_or_url or "").strip().strip('"\'')
    if not origin:
        return SourceText(False, "", "error", origin, "No recibí ninguna fuente.")

    low = origin.lower()
    # Página web normal (http/https) que NO es un archivo descargable: se lee el HTML.
    if low.startswith(("http://", "https://")) and os.path.splitext(low.split("?")[0])[1] not in (
        ".pdf", ".docx", ".pptx", ".epub", ".txt", ".md"):
        return _read_web(origin, max_chars)

    # file:// o http a un archivo -> obtenemos una ruta local real (descargando si hace falta).
    local = _ensure_local_file(origin)
    if not local or not os.path.exists(local):
        return SourceText(False, "", "error", origin,
                          "No encuentro ese archivo o no pude abrir el enlace.")

    ext = os.path.splitext(local)[1].lower()
    if ext in (".txt", ".md"):
        return _read_txt(local, max_chars)
    if ext == ".pdf":
        return _read_pdf(local, max_chars)
    if ext == ".docx":
        return _read_docx(local, max_chars)
    if ext == ".pptx":
        return _read_pptx(local, max_chars)
    if ext == ".epub":
        return _read_epub(local, max_chars)
    if ext in (".png", ".jpg", ".jpeg", ".bmp", ".webp"):
        return _read_ocr(local, max_chars)
    return SourceText(False, "", "error", origin, f"No sé leer archivos {ext} todavía.")


# ---------------------------------------------------------------- txt
def _read_txt(path: str, max_chars: int) -> SourceText:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return SourceText(True, f.read()[:max_chars], "txt", path)
    except Exception as exc:
        return SourceText(False, "", "error", path, f"Error leyendo el texto: {exc}")


# ---------------------------------------------------------------- pdf
def _read_pdf(path: str, max_chars: int) -> SourceText:
    # Lectura PÁGINA POR PÁGINA: usa el texto embebido donde lo hay y hace OCR solo
    # en las páginas escaneadas. Así funciona con PDFs de texto, escaneados o MIXTOS.
    # (Si no hay motor de render/OCR, degrada a extraer solo el texto embebido.)
    texto = ""
    try:
        texto = read_pdf_full_deep(path)
    except Exception as exc:
        print("[profesora] lectura profunda del PDF falló:", exc)

    # Respaldo: extracción simple clásica si lo anterior no consiguió nada.
    if len((texto or "").strip()) < _MIN_TEXT_CHARS:
        simple = ""
        try:
            import pdfplumber
            partes = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    partes.append(page.extract_text() or "")
                    if sum(len(p) for p in partes) > max_chars:
                        break
            simple = "\n".join(partes)
        except Exception:
            for modname in ("pypdf", "PyPDF2"):
                try:
                    mod = __import__(modname)
                    reader = mod.PdfReader(path)
                    partes = []
                    for page in reader.pages:
                        partes.append(page.extract_text() or "")
                        if sum(len(p) for p in partes) > max_chars:
                            break
                    simple = "\n".join(partes)
                    break
                except Exception:
                    continue
        if len((simple or "").strip()) > len((texto or "").strip()):
            texto = simple

    if (texto or "").strip():
        return SourceText(True, texto[:max_chars], "pdf", path)
    return SourceText(False, "", "error", path,
                      "No pude extraer nada del PDF. Para PDFs escaneados instala "
                      "PyMuPDF y Tesseract (pip install pymupdf pytesseract).")


# ---------------------------------------------------------------- docx
def _read_docx(path: str, max_chars: int) -> SourceText:
    try:
        import docx
        doc = docx.Document(path)
        texto = "\n".join(p.text for p in doc.paragraphs)
        return SourceText(True, texto[:max_chars], "docx", path)
    except Exception:
        return SourceText(False, "", "error", path,
                          "Para leer Word instala python-docx (pip install python-docx).")


# ---------------------------------------------------------------- pptx
def _read_pptx(path: str, max_chars: int) -> SourceText:
    try:
        from pptx import Presentation
        prs = Presentation(path)
        partes = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    partes.append(shape.text_frame.text)
        return SourceText(True, "\n".join(partes)[:max_chars], "pptx", path)
    except Exception:
        return SourceText(False, "", "error", path,
                          "Para leer PowerPoint instala python-pptx (pip install python-pptx).")


# ---------------------------------------------------------------- epub
def _read_epub(path: str, max_chars: int) -> SourceText:
    try:
        from ebooklib import epub, ITEM_DOCUMENT
        libro = epub.read_epub(path)
        partes = []
        for item in libro.get_items_of_type(ITEM_DOCUMENT):
            partes.append(_strip_html(item.get_content().decode("utf-8", "ignore")))
            if sum(len(p) for p in partes) > max_chars:
                break
        return SourceText(True, "\n".join(partes)[:max_chars], "epub", path)
    except Exception:
        # Respaldo sin librerías: un EPUB es un ZIP con XHTML dentro.
        try:
            partes = []
            with zipfile.ZipFile(path) as z:
                for name in z.namelist():
                    if name.lower().endswith((".xhtml", ".html", ".htm")):
                        partes.append(_strip_html(z.read(name).decode("utf-8", "ignore")))
                        if sum(len(p) for p in partes) > max_chars:
                            break
            if partes:
                return SourceText(True, "\n".join(partes)[:max_chars], "epub", path,
                                  "Leído sin ebooklib (respaldo por ZIP).")
        except Exception as exc:
            return SourceText(False, "", "error", path, f"No pude abrir el EPUB: {exc}")
    return SourceText(False, "", "error", path,
                      "Para leer EPUB instala ebooklib (pip install EbookLib).")


# ---------------------------------------------------------------- ocr
def _read_ocr(path: str, max_chars: int) -> SourceText:
    # 1) Reusar el OCR que ya tiene YUE, si expone una función utilizable.
    try:
        from core import screen_text
        for fn in ("ocr_image", "image_to_text", "read_image", "extract_text"):
            f = getattr(screen_text, fn, None)
            if callable(f):
                txt = f(path)
                if txt:
                    return SourceText(True, str(txt)[:max_chars], "ocr", path,
                                      "OCR con el motor propio de YUE.")
    except Exception:
        pass
    # 2) Respaldo: pytesseract.
    try:
        import pytesseract
        from PIL import Image
        txt = pytesseract.image_to_string(Image.open(path), lang="spa+eng")
        return SourceText(True, txt[:max_chars], "ocr", path)
    except Exception:
        return SourceText(False, "", "error", path,
                          "Para OCR de imágenes instala pytesseract y Tesseract.")


# ---------------------------------------------------------------- web
def _read_web(url: str, max_chars: int) -> SourceText:
    try:
        import urllib.request
        req = urllib.request.Request(url, headers={"User-Agent": "YUE-Teacher/1.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            raw = resp.read().decode("utf-8", "ignore")
        return SourceText(True, _strip_html(raw)[:max_chars], "web", url)
    except Exception as exc:
        return SourceText(False, "", "error", url, f"No pude descargar la página: {exc}")


def _strip_html(html: str) -> str:
    html = re.sub(r"(?is)<(script|style|head).*?</\1>", " ", html)
    html = re.sub(r"(?s)<[^>]+>", " ", html)
    html = re.sub(r"&[a-z]+;", " ", html)
    return re.sub(r"\s+", " ", html).strip()


def read_pdf_pages(path: str, max_pages: int = 60, max_chars_per_page: int = 4000) -> list:
    """Devuelve una lista con el TEXTO de cada página del PDF (para lectura guiada).

    Usa pdfplumber si está; si no, pypdf/PyPDF2. Si no hay ninguno, devuelve [].
    Cada elemento es el texto de una página (puede ir vacío si la página es imagen).
    """
    paginas: list = []
    # NUEVO: acepta file:///... y enlaces http; los convierte a un archivo local.
    local = _ensure_local_file(path)
    if not local or not os.path.exists(local):
        return paginas
    try:
        import pdfplumber
        with pdfplumber.open(local) as pdf:
            for i, page in enumerate(pdf.pages):
                if i >= max_pages:
                    break
                paginas.append((page.extract_text() or "")[:max_chars_per_page])
        return paginas
    except Exception:
        pass
    for modname in ("pypdf", "PyPDF2"):
        try:
            mod = __import__(modname)
            reader = mod.PdfReader(local)
            for i, page in enumerate(reader.pages):
                if i >= max_pages:
                    break
                paginas.append((page.extract_text() or "")[:max_chars_per_page])
            return paginas
        except Exception:
            continue
    return paginas


# ==========================================================================
# Lectura PROFUNDA de PDF por página: texto embebido -> OCR -> (marca visión).
# Permite leer PDFs escaneados, mixtos o de puras imágenes. Es LAZY: se procesa
# una página cuando toca explicarla, para no hacer OCR de todo el documento a la vez.
# ==========================================================================
_MIN_TEXT_CHARS = 40   # a partir de aquí consideramos que la página ya trae texto útil


def page_count(path_or_url: str) -> int:
    """Número de páginas del PDF (funciona también con PDFs escaneados)."""
    local = _ensure_local_file(path_or_url)
    if not local or not os.path.exists(local):
        return 0
    try:
        import fitz  # PyMuPDF
        with fitz.open(local) as doc:
            return int(doc.page_count)
    except Exception:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(local) as pdf:
            return len(pdf.pages)
    except Exception:
        pass
    for modname in ("pypdf", "PyPDF2"):
        try:
            mod = __import__(modname)
            return len(mod.PdfReader(local).pages)
        except Exception:
            continue
    return 0


def _page_embedded_text(local: str, index: int, max_chars: int = 4000) -> str:
    """Texto EMBEBIDO de una página (sin OCR). Vacío si la página es imagen."""
    try:
        import fitz
        with fitz.open(local) as doc:
            if 0 <= index < doc.page_count:
                return (doc.load_page(index).get_text() or "")[:max_chars]
    except Exception:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(local) as pdf:
            if 0 <= index < len(pdf.pages):
                return (pdf.pages[index].extract_text() or "")[:max_chars]
    except Exception:
        pass
    for modname in ("pypdf", "PyPDF2"):
        try:
            mod = __import__(modname)
            reader = mod.PdfReader(local)
            if 0 <= index < len(reader.pages):
                return (reader.pages[index].extract_text() or "")[:max_chars]
        except Exception:
            continue
    return ""


def render_pdf_page_png(path_or_url: str, index: int, dpi: int = 180) -> str | None:
    """Rasteriza UNA página del PDF a un PNG temporal. Prioridad: PyMuPDF (sin
    dependencias del sistema) y, si no está, pdf2image (requiere Poppler).
    Devuelve la ruta del PNG (el llamador debe borrarlo al terminar) o None."""
    local = _ensure_local_file(path_or_url)
    if not local or not os.path.exists(local):
        return None
    import tempfile
    # 1) PyMuPDF (fitz): recomendado para el .exe (un solo wheel, sin binarios).
    try:
        import fitz
        with fitz.open(local) as doc:
            if not (0 <= index < doc.page_count):
                return None
            zoom = max(1.0, dpi / 72.0)
            pix = doc.load_page(index).get_pixmap(matrix=fitz.Matrix(zoom, zoom))
            fd, tmp = tempfile.mkstemp(suffix=".png")
            os.close(fd)
            pix.save(tmp)
            return tmp
    except Exception:
        pass
    # 2) pdf2image (necesita Poppler instalado en el sistema).
    try:
        from pdf2image import convert_from_path
        imgs = convert_from_path(local, dpi=dpi, first_page=index + 1, last_page=index + 1)
        if imgs:
            fd, tmp = tempfile.mkstemp(suffix=".png")
            os.close(fd)
            imgs[0].save(tmp, "PNG")
            return tmp
    except Exception as exc:
        print("[profesora] no pude rasterizar la página del PDF:", exc)
    return None


def ocr_image_text(png_path: str, max_chars: int = 4000) -> str:
    """OCR de una imagen reutilizando el motor OCR de YUE (pytesseract de respaldo)."""
    if not png_path:
        return ""
    src = _read_ocr(png_path, max_chars)
    return src.text if src.ok else ""


def read_pdf_page_deep(path_or_url: str, index: int, dpi: int = 180,
                       do_ocr: bool = True, max_chars: int = 4000) -> dict:
    """Lee una página con capas y dice cómo debe explicarse.

    Devuelve un dict:
        text         -> texto disponible (embebido y/o OCR)
        image        -> ruta PNG de la página rasterizada (o "")
        source       -> "texto" | "ocr" | "imagen"
        needs_vision -> True si es una página de imagen/diagrama y conviene que
                        YUE la MIRE con su visión para explicarla como profesora.
    """
    local = _ensure_local_file(path_or_url) or path_or_url
    texto = _page_embedded_text(local, index, max_chars)
    if len((texto or "").strip()) >= _MIN_TEXT_CHARS:
        return {"text": texto, "image": "", "source": "texto", "needs_vision": False}

    png = render_pdf_page_png(local, index, dpi) if do_ocr else None
    ocr_txt = ocr_image_text(png, max_chars) if png else ""
    combinado = ((texto or "") + "\n" + (ocr_txt or "")).strip()

    if len((ocr_txt or "").strip()) >= _MIN_TEXT_CHARS:
        # Página escaneada de texto: el OCR bastó para leerla.
        return {"text": combinado, "image": png or "", "source": "ocr", "needs_vision": False}
    # Ni texto ni OCR útil: es imagen/diagrama -> que YUE la observe con su visión.
    return {"text": combinado, "image": png or "", "source": "imagen", "needs_vision": True}


def read_pdf_full_deep(path_or_url: str, max_pages: int = 40, dpi: int = 160) -> str:
    """Escaneo COMPLETO de un PDF (texto + OCR de páginas escaneadas), concatenado.

    Útil cuando se quiere todo el contenido de golpe (resumen, examen, etc.).
    Las páginas de pura imagen aportan lo que el OCR consiga; para explicarlas por
    su contenido visual se usa la lectura guiada página por página con visión.
    """
    local = _ensure_local_file(path_or_url)
    if not local or not os.path.exists(local):
        return ""
    total = min(page_count(local), max_pages)
    trozos = []
    pngs = []
    try:
        for i in range(total):
            info = read_pdf_page_deep(local, i, dpi=dpi)
            if info.get("image"):
                pngs.append(info["image"])
            t = (info.get("text") or "").strip()
            if t:
                trozos.append(f"[Página {i + 1}]\n{t}")
    finally:
        for p in pngs:
            try:
                os.remove(p)
            except Exception:
                pass
    return "\n\n".join(trozos)


def capabilities() -> dict:
    """Qué formatos puede leer ahora mismo (para informar al usuario/UI)."""
    def _tiene(mod):
        try:
            __import__(mod)
            return True
        except Exception:
            return False
    return {
        "txt": True,
        "web": True,
        "pdf": _tiene("pdfplumber") or _tiene("pypdf") or _tiene("PyPDF2") or _tiene("fitz"),
        "pdf_escaneado": (_tiene("fitz") or _tiene("pdf2image")) and _tiene("pytesseract"),
        "render": _tiene("fitz") or _tiene("pdf2image"),
        "docx": _tiene("docx"),
        "pptx": _tiene("pptx"),
        "epub": _tiene("ebooklib") or True,   # respaldo ZIP siempre disponible
        "ocr": _tiene("pytesseract"),
    }
