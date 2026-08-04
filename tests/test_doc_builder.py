"""Pruebas del generador de trabajos con diseño (Excel + PowerPoint).

    python -m pytest tests/test_doc_builder.py
    python tests/test_doc_builder.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.doc_builder import (
    TablePlan, DeckPlan, get_palette, PALETTES,
    parse_table_plan, parse_deck_plan, build_table, build_deck,
)


def test_paletas_existen():
    assert "colegio" in PALETTES
    p = get_palette("colegio")
    assert len(p.primary) == 6  # hex sin '#'


def test_parse_table_plan_desde_json():
    js = '{"title":"Notas","headers":["Alumno","Nota"],"rows":[["Ana","18"]]}'
    plan = parse_table_plan(js)
    assert plan.title == "Notas"
    assert plan.headers == ["Alumno", "Nota"]
    assert plan.rows == [["Ana", "18"]]


def test_parse_quita_fences_y_texto():
    js = 'Claro, aquí tienes:\n```json\n{"title":"X","headers":["A"],"rows":[["1"]]}\n```'
    plan = parse_table_plan(js)
    assert plan.title == "X"


def test_table_plan_normaliza_filas():
    plan = TablePlan(title="T", headers=["A", "B", "C"],
                     rows=[["1"], ["1", "2", "3", "4"]]).validate()
    assert all(len(r) == 3 for r in plan.rows)  # rellena/recorta a 3 columnas


def test_table_plan_valida_sin_titulo():
    try:
        TablePlan(title="", headers=["A"], rows=[]).validate()
        assert False, "debió fallar sin título"
    except ValueError:
        pass


def test_build_table_genera_xlsx_real():
    from openpyxl import load_workbook
    d = tempfile.mkdtemp(prefix="yue_xlsx_")
    out = str(Path(d) / "notas.xlsx")
    plan = TablePlan(
        title="Notas del 3er grado",
        subtitle="Bimestre I",
        sheet_name="Notas",
        palette="colegio",
        headers=["Alumno", "Matemática", "Comunicación"],
        rows=[["Ana", 18, 17], ["Luis", 15, 16], ["Marta", 20, 19]],
        total_columns=[1, 2],
        note="Fuente: registro del docente",
    )
    ruta = build_table(plan, out)
    assert Path(ruta).exists()
    wb = load_workbook(ruta)
    ws = wb["Notas"]
    # El título está en la primera fila.
    assert "Notas del 3er grado" in str(ws.cell(row=1, column=1).value)
    # Hay una fila de TOTAL con fórmula SUMA en columnas numéricas.
    encontrado_total = False
    for row in ws.iter_rows():
        for c in row:
            if c.value == "TOTAL":
                encontrado_total = True
            if isinstance(c.value, str) and c.value.startswith("=SUM("):
                assert True
    assert encontrado_total


def test_parse_deck_plan_desde_json():
    js = ('{"title":"La fotosíntesis","palette":"verde","slides":['
          '{"type":"title","title":"La fotosíntesis"},'
          '{"type":"bullets","title":"Qué es","bullets":["Transforma luz","Da oxígeno"]}]}')
    deck = parse_deck_plan(js)
    assert deck.title == "La fotosíntesis"
    assert deck.palette == "verde"
    assert len(deck.slides) == 2


def test_deck_plan_tipo_desconocido_cae_a_bullets():
    deck = DeckPlan(title="T", slides=[{"type": "inventado", "title": "X"}]).validate()
    assert deck.slides[0]["type"] == "bullets"


def test_build_deck_genera_pptx_real():
    from pptx import Presentation
    d = tempfile.mkdtemp(prefix="yue_pptx_")
    out = str(Path(d) / "clase.pptx")
    deck = DeckPlan(
        title="La célula",
        subtitle="Ciencias Naturales",
        author="Prof. YUE",
        palette="morado_yue",
        slides=[
            {"type": "title", "title": "La célula", "subtitle": "Ciencias"},
            {"type": "section", "title": "Partes de la célula"},
            {"type": "bullets", "title": "Componentes",
             "bullets": ["Núcleo", "Membrana", "Citoplasma"]},
            {"type": "table", "title": "Comparación",
             "headers": ["Tipo", "Pared"], "rows": [["Animal", "No"], ["Vegetal", "Sí"]]},
            {"type": "closing", "title": "¡Gracias!", "subtitle": "A estudiar"},
        ],
    )
    ruta = build_deck(deck, out)
    assert Path(ruta).exists()
    prs = Presentation(ruta)
    assert len(prs.slides) == 5  # las 5 diapositivas del plan


def test_build_deck_agrega_portada_si_falta():
    from pptx import Presentation
    d = tempfile.mkdtemp(prefix="yue_pptx2_")
    out = str(Path(d) / "sinportada.pptx")
    deck = DeckPlan(title="Tema", slides=[
        {"type": "bullets", "title": "Punto", "bullets": ["a", "b"]},
    ])
    build_deck(deck, out)
    prs = Presentation(out)
    # Se antepone la portada -> 2 diapositivas (portada + la de viñetas).
    assert len(prs.slides) == 2


if __name__ == "__main__":
    fallos = []
    for _n, _f in sorted(globals().items()):
        if _n.startswith("test_") and callable(_f):
            try:
                _f()
                print("  OK  " + _n)
            except Exception as _e:
                print("  FALLO  " + _n + f"  ·  {_e}")
                fallos.append(_n)
    print("\n" + ("TODO OK" if not fallos else f"FALLOS: {fallos}"))
    sys.exit(1 if fallos else 0)
